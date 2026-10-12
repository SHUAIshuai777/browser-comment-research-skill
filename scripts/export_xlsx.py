"""Optional Excel export; IDs and text stored as strings, Beijing wall-clock dates numeric."""
import argparse,json
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook,load_workbook
from openpyxl.utils.datetime import to_excel
ap=argparse.ArgumentParser();ap.add_argument('--project',type=Path,required=True);a=ap.parse_args();p=a.project
wb=Workbook();ws=wb.active;ws.title='评论';ws.append(['任务','评论身份','评论ID','平台','视频ID','原始正文','北京时间','材料类型','样本组','人工编码'])
for result in json.loads((p/'outputs/result.json').read_text()):
 ks={m['comment_key'] for m in result['sample_memberships'] if m['layer']=='K'}
 for r in result['root_records']:
  ws.append([r['unit_id'],r['comment_key'],r['comment_id'],r['platform'],r['video_id'],r['raw_text'],None,r['content_kind'],'K' if r['comment_key'] in ks else 'neutral',''])
  for col in [1,2,3,4,5,6,8,9,10]:ws.cell(ws.max_row,col).data_type='s'
  if r.get('time_lower'):
   dt=datetime.fromisoformat(r['time_lower']);assert dt.utcoffset().total_seconds()==28800
   c=ws.cell(ws.max_row,7,to_excel(dt.replace(tzinfo=None)));c.number_format='yyyy-mm-dd hh:mm:ss'
ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
for key,width in {'A':24,'B':55,'C':25,'D':12,'E':25,'F':80,'G':25,'H':16,'I':15,'J':20}.items():ws.column_dimensions[key].width=width
from openpyxl.styles import Alignment
for row in ws:
 for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
audit=wb.create_sheet('任务审计');audit.append(['任务','状态','缺口与排序提示'])
for row in json.loads((p/'outputs/audit.json').read_text()):audit.append([row['unit_id'],row['status'],json.dumps(row,ensure_ascii=False)])
dest=p/'outputs/comments.xlsx';wb.save(dest)
check=load_workbook(dest)
assert not any(c.data_type=='f' for sheet in check for row in sheet for c in row)
print(dest)

import unittest,tempfile,subprocess,sys,json
from pathlib import Path
from datetime import datetime
ROOT=Path(__file__).resolve().parents[1]
try:
 from openpyxl import load_workbook
except ImportError:load_workbook=None
@unittest.skipUnless(load_workbook,'optional openpyxl not installed')
class ExcelTest(unittest.TestCase):
 def test_ids_dates_and_formula_text(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'demo';subprocess.run([sys.executable,str(ROOT/'scripts/research.py'),'demo','--project',str(p)],check=True,stdout=subprocess.DEVNULL)
   result=json.loads((p/'outputs/result.json').read_text());result[0]['root_records'][0]['raw_text']='=1+1';(p/'outputs/result.json').write_text(json.dumps(result))
   subprocess.run([sys.executable,str(ROOT/'scripts/export_xlsx.py'),'--project',str(p)],check=True,stdout=subprocess.DEVNULL)
   sheet=load_workbook(p/'outputs/comments.xlsx')['评论']
   for row,r in zip(list(sheet)[1:],[r for unit in result for r in unit['root_records']]):
    self.assertEqual(row[2].value,r['comment_id']);self.assertEqual(row[2].data_type,'s');self.assertEqual(row[5].value,r['raw_text']);self.assertNotEqual(row[5].data_type,'f');self.assertEqual(row[6].value,datetime.fromisoformat(r['time_lower']).replace(tzinfo=None))
